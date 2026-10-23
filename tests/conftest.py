from decimal import Decimal

import pytest
from django.core.files import File
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook

from apps.academics.models import AcademicYear, ClassRoom, GradeRecord, Student
from apps.accounts.models import User
from apps.core.text import normalize_label, normalize_year_label
from apps.imports.models import SourceFile

HEADERS = ['Nom complet', 'Pourcentage', 'Classe', 'Section', 'Année scolaire']
PASSWORD = 'motdepasse-solide-1'


# ── Utilisateurs ─────────────────────────────────────────────────

@pytest.fixture
def make_user(db):
    def _make(username, role):
        return User.objects.create_user(username=username, password=PASSWORD, role=role)
    return _make


@pytest.fixture
def admin(make_user):
    return make_user('admin1', User.ROLE_ADMIN)


@pytest.fixture
def editor(make_user):
    return make_user('editor1', User.ROLE_EDITOR)


@pytest.fixture
def editor2(make_user):
    return make_user('editor2', User.ROLE_EDITOR)


@pytest.fixture
def reader(make_user):
    return make_user('reader1', User.ROLE_READER)


@pytest.fixture
def as_user(client):
    def _login(user):
        client.force_login(user)
        return client
    return _login


# ── Données ──────────────────────────────────────────────────────

@pytest.fixture
def make_grade(db):
    """Crée un résultat (et l'élève, la classe, l'année s'ils n'existent pas)."""
    def _make(name='JEAN DUPONT', percentage=60, classroom='7 EB A', section='SECONDAIRE',
              year='2024-2025', **extra):
        # Les modèles stockent tout en majuscules : on cherche avec la forme normalisée
        student, _ = Student.objects.get_or_create(full_name=normalize_label(name))
        room, _ = ClassRoom.objects.get_or_create(name=normalize_label(classroom), section=normalize_label(section))
        academic_year, _ = AcademicYear.objects.get_or_create(label=normalize_year_label(year))
        value = None if percentage is None else Decimal(str(percentage))
        return GradeRecord.objects.create(
            student=student, classroom=room, academic_year=academic_year, percentage=value, **extra
        )
    return _make


# ── Fichiers Excel ───────────────────────────────────────────────

@pytest.fixture
def make_xlsx(tmp_path):
    counter = iter(range(1, 1000))

    def _make(rows, *, headers=HEADERS, extra_sheets=None, pct_format=False, title_rows=0):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'Feuil1'
        for index in range(title_rows):
            sheet.append([f'Titre {index}'])
        if headers is not None:
            sheet.append(headers)
        for row in rows:
            sheet.append(list(row))
        if pct_format:
            for cell in sheet['B'][title_rows + 1:]:
                cell.number_format = '0.0%'
        for name, sheet_rows in (extra_sheets or {}).items():
            extra = workbook.create_sheet(name)
            for row in sheet_rows:
                extra.append(list(row))
        path = tmp_path / f'palmares-{next(counter)}.xlsx'
        workbook.save(path)
        return str(path)
    return _make


@pytest.fixture
def make_xls(tmp_path):
    xlwt = pytest.importorskip('xlwt')

    def _make(rows, *, percent_rows=False):
        workbook = xlwt.Workbook()
        sheet = workbook.add_sheet('Feuil1')
        for col, header in enumerate(HEADERS):
            sheet.write(0, col, header)
        percent_style = xlwt.easyxf(num_format_str='0.0%')
        for r, row in enumerate(rows, start=1):
            for col, value in enumerate(row):
                if col == 1 and percent_rows and isinstance(value, (int, float)):
                    sheet.write(r, col, value, percent_style)
                else:
                    sheet.write(r, col, value)
        path = tmp_path / 'ancien.xls'
        workbook.save(str(path))
        return str(path)
    return _make


@pytest.fixture
def source_file(db):
    """Enregistre un fichier déjà écrit sur disque comme SourceFile (sans lancer d'analyse)."""
    def _make(user, path, name='palmares.xlsx', **fields):
        with open(path, 'rb') as handle:
            return SourceFile.objects.create(
                imported_by=user, original_filename=name, file=File(handle, name=name), **fields
            )
    return _make


@pytest.fixture
def upload():
    """Fichier prêt à être envoyé dans un formulaire."""
    def _upload(path, name='palmares.xlsx'):
        with open(path, 'rb') as handle:
            return SimpleUploadedFile(name, handle.read())
    return _upload
