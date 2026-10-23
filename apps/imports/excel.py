"""Lecture des classeurs Excel (.xlsx et .xls) et validation des lignes.

Aucune base de données ici : le module ne fait que transformer un fichier en lignes propres.
"""
import datetime
import logging
import zipfile
from collections import namedtuple
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP

from apps.core.text import collapse_spaces, name_key, normalize_label, normalize_year_label

logger = logging.getLogger(__name__)

EXPECTED_COLUMNS = ('Nom complet', 'Pourcentage', 'Classe', 'Section', 'Année scolaire')
_WANTED = {name_key(column): column for column in EXPECTED_COLUMNS}

MAX_ROWS = 50_000                       # lignes de données par fichier
MAX_UNCOMPRESSED_BYTES = 300 * 1024 * 1024   # protection contre les « zip bombs »
HEADER_SCAN_ROWS = 10                   # l'en-tête peut être précédé de lignes de titre
BLANK_ROWS_BEFORE_STOP = 500            # arrêt de la lecture après N lignes vides consécutives

MAX_LENGTHS = {'Nom complet': 200, 'Classe': 100, 'Section': 100, 'Année scolaire': 20}
_EMPTY_MARKERS = {'', 'n/a', 'na', '-', '--', '/', '—'}
_CENTS = Decimal('0.01')

Cell = namedtuple('Cell', 'value percent')   # percent : la cellule est au format « % » d'Excel
_EMPTY = Cell(None, False)

XLSX_MAGIC = b'PK\x03\x04'
XLS_MAGIC = b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1'


class ImportFileError(Exception):
    """Fichier illisible ou structure invalide. Le message est destiné à l'utilisateur."""

    def __init__(self, message, missing_columns=()):
        super().__init__(message)
        self.missing_columns = list(missing_columns)


class Sheet:
    def __init__(self, name, rows, fraction_scaled):
        self.name = name
        self.rows = rows                        # [(numéro de ligne, {colonne: Cell})]
        self.fraction_scaled = fraction_scaled  # pourcentages saisis en fractions (0,75 = 75 %)


# ── Détection du format ──────────────────────────────────────────

def detect_format(path):
    """Se base sur le contenu réel du fichier, pas sur son extension."""
    with open(path, 'rb') as handle:
        head = handle.read(8)
    if head.startswith(XLSX_MAGIC):
        return 'xlsx'
    if head.startswith(XLS_MAGIC):
        return 'xls'
    raise ImportFileError("Le fichier n'est pas un classeur Excel valide (.xlsx ou .xls).")


def looks_like_excel(head):
    return head.startswith(XLSX_MAGIC) or head.startswith(XLS_MAGIC)


# ── Conversion des cellules ──────────────────────────────────────

def cell_text(value):
    if value is None:
        return ''
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else repr(value)
    if isinstance(value, (datetime.datetime, datetime.date)):
        return value.isoformat()
    return str(value)


def parse_percentage(raw, percent_format=False, fraction_scaled=False):
    """Convertit une saisie en pourcentage. Retourne (Decimal | None, message d'erreur | None).

    Accepte « 75 », « 75,5 », « 75 % », les nombres Excel et les cellules au format %.
    Une valeur vide (ou « N/A », « - », « / ») donne (None, None) : élève sans note.
    """
    if raw is None:
        return None, None
    if isinstance(raw, bool):
        return None, f"Pourcentage invalide ({raw!r})"

    scale = percent_format or fraction_scaled
    try:
        if isinstance(raw, (int, float)):
            value = Decimal(repr(raw))
        else:
            text = collapse_spaces(raw)
            if text.lower() in _EMPTY_MARKERS:
                return None, None
            if text.endswith('%'):
                text, scale = text[:-1].strip(), False
            value = Decimal(text.replace(',', '.').replace(' ', ''))
    except InvalidOperation:
        return None, f"Pourcentage invalide ({collapse_spaces(raw)!r})"

    if not value.is_finite():
        return None, f"Pourcentage invalide ({collapse_spaces(raw)!r})"
    if scale:
        value *= 100
    # Contrôle AVANT d'arrondir : quantize() lève une exception sur les valeurs gigantesques
    if not (Decimal('0') <= value <= Decimal('100')):
        return None, f"Pourcentage hors limites ({value:.6G})"
    return value.quantize(_CENTS, rounding=ROUND_HALF_UP), None


def validate_row(cells, fraction_scaled=False):
    """Valide et nettoie une ligne. Retourne (valeurs nettoyées, liste d'erreurs)."""
    errors = []
    cleaned = {}

    full_name = normalize_label(cell_text(cells['Nom complet'].value))
    classroom = normalize_label(cell_text(cells['Classe'].value))
    section = normalize_label(cell_text(cells['Section'].value))
    year = normalize_year_label(cell_text(cells['Année scolaire'].value))

    for column, value in (('Nom complet', full_name), ('Classe', classroom),
                          ('Section', section), ('Année scolaire', year)):
        if not value:
            errors.append(f"{column} manquant(e)")
        elif len(value) > MAX_LENGTHS[column]:
            errors.append(f"{column} trop long(ue) (max {MAX_LENGTHS[column]} caractères)")

    cell = cells['Pourcentage']
    percentage, error = parse_percentage(cell.value, cell.percent, fraction_scaled)
    if error:
        errors.append(error)

    cleaned.update(
        full_name=full_name, classroom_name=classroom, section=section,
        year_label=year, percentage=percentage,
    )
    return cleaned, errors


# ── Lecture des feuilles ─────────────────────────────────────────

def _detect_fractions(rows):
    """Vrai si tous les pourcentages sont entre 0 et 1 (colonne saisie en fractions)."""
    highest = None
    for _, cells in rows:
        cell = cells['Pourcentage']
        if cell.percent:
            return False
        value, error = parse_percentage(cell.value)
        if error or value is None:
            continue
        if highest is None or value > highest:
            highest = value
    return highest is not None and Decimal('0') < highest <= Decimal('1')


def _collect(name, rows_factory):
    """Cherche l'en-tête puis lit les lignes de données.
    Retourne (Sheet | None, colonnes manquantes)."""
    header_row = None
    mapping = {}
    best_missing = list(EXPECTED_COLUMNS)
    for index, row in enumerate(rows_factory(1, HEADER_SCAN_ROWS, 100)):
        found = {}
        for position, cell in enumerate(row):
            column = _WANTED.get(name_key(cell_text(cell.value)))
            if column and column not in found:
                found[column] = position
        missing = [column for column in EXPECTED_COLUMNS if column not in found]
        if not missing:
            header_row, mapping = index, found
            break
        if index == 0 or len(missing) < len(best_missing):
            best_missing = missing
    if header_row is None:
        return None, best_missing

    width = max(mapping.values()) + 1
    first_row = header_row + 2   # numéro (à partir de 1) de la première ligne de données
    rows = []
    blank_streak = 0
    for number, row in enumerate(rows_factory(first_row, None, width), start=first_row):
        cells = {column: (row[position] if position < len(row) else _EMPTY)
                 for column, position in mapping.items()}
        if all(not cell_text(cell.value).strip() for cell in cells.values()):
            blank_streak += 1
            if blank_streak >= BLANK_ROWS_BEFORE_STOP:
                break
            continue
        blank_streak = 0
        rows.append((number, cells))
        if len(rows) > MAX_ROWS:
            raise ImportFileError(f"Fichier trop volumineux : plus de {MAX_ROWS} lignes.")
    return Sheet(name, rows, _detect_fractions(rows)), []


def _read_xlsx(path):
    try:
        with zipfile.ZipFile(path) as archive:
            if sum(info.file_size for info in archive.infolist()) > MAX_UNCOMPRESSED_BYTES:
                raise ImportFileError("Fichier trop volumineux une fois décompressé.")
    except zipfile.BadZipFile as exc:
        raise ImportFileError("Le fichier .xlsx est corrompu.") from exc

    import openpyxl
    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        logger.warning("Lecture xlsx impossible : %s", exc)
        raise ImportFileError("Impossible de lire le fichier Excel (.xlsx).") from exc

    def sheet_reader(worksheet):
        def rows(min_row, max_row, max_col):
            for row in worksheet.iter_rows(min_row=min_row, max_row=max_row, max_col=max_col):
                yield [Cell(c.value, '%' in (c.number_format or '')) for c in row]
        return rows

    try:
        return [(ws.title, sheet_reader(ws)) for ws in workbook.worksheets], workbook.close
    except Exception:
        workbook.close()
        raise


def _read_xls(path):
    import xlrd
    try:
        book = xlrd.open_workbook(path, formatting_info=True, on_demand=True)
    except Exception as exc:
        logger.warning("Lecture xls impossible : %s", exc)
        raise ImportFileError("Impossible de lire le fichier Excel (.xls).") from exc

    def convert(cell):
        kind = cell.ctype
        if kind in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
            return _EMPTY
        if kind == xlrd.XL_CELL_ERROR:
            return Cell('#ERREUR', False)
        value = cell.value
        percent = False
        if kind == xlrd.XL_CELL_NUMBER:
            fmt = book.format_map[book.xf_list[cell.xf_index].format_key].format_str
            percent = '%' in fmt
        elif kind == xlrd.XL_CELL_DATE:
            try:
                value = xlrd.xldate_as_datetime(value, book.datemode)
            except Exception:
                value = str(value)
        elif kind == xlrd.XL_CELL_BOOLEAN:
            value = bool(value)
        return Cell(value, percent)

    def sheet_reader(sheet):
        def rows(min_row, max_row, max_col):
            last = sheet.nrows if max_row is None else min(max_row, sheet.nrows)
            for r in range(min_row - 1, last):
                yield [convert(sheet.cell(r, c)) for c in range(min(max_col, sheet.ncols))]
        return rows

    sheets = [(book.sheet_names()[i], sheet_reader(book.sheet_by_index(i))) for i in range(book.nsheets)]
    return sheets, book.release_resources


def read_workbook(path):
    """Lit toutes les feuilles contenant les colonnes attendues.
    Retourne (feuilles, [(feuille ignorée, colonnes manquantes)])."""
    kind = detect_format(path)
    readers, close = _read_xlsx(path) if kind == 'xlsx' else _read_xls(path)
    sheets, ignored = [], []
    try:
        for name, rows_factory in readers:
            sheet, missing = _collect(name, rows_factory)
            if sheet is None:
                ignored.append((name, missing))
            elif sheet.rows:
                sheets.append(sheet)
    finally:
        close()

    if not sheets:
        if ignored and len(ignored) == len(readers):
            missing = ignored[0][1]
            raise ImportFileError(f"Colonnes manquantes : {', '.join(missing)}", missing_columns=missing)
        raise ImportFileError("Le fichier ne contient aucune ligne de données.")
    return sheets, ignored
